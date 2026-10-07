"""
解析微信支付导出的 Excel 账单（.xlsx），提取每笔交易核心字段。

依赖: openpyxl
安装: pip install -r requirements.txt
"""

from __future__ import annotations

import argparse
import json
from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any, Iterator

from openpyxl import load_workbook

# Excel 明细行共 11 列，解析时只取前几列
_SHEET_COLUMN_COUNT = 11

HEADER_MARKER = "交易时间"


@dataclass
class Transaction:
    """单笔交易（仅保留常用字段）。"""

    交易时间: str  # 日期，格式 YYYY-MM-DD
    交易类型: str
    交易对方: str
    收_支: str
    金额: float | None

    def to_dict(self) -> dict[str, Any]:
        return {
            "交易时间": self.交易时间,
            "交易类型": self.交易类型,
            "交易对方": self.交易对方,
            "收/支": self.收_支,
            "金额": self.金额,
        }


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _excel_date(value: Any) -> str:
    """将单元格值转为日期字符串 YYYY-MM-DD。"""
    if value is None or value == "":
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, (int, float)):
        base = datetime(1899, 12, 30)
        dt = base + timedelta(days=float(value))
        return dt.strftime("%Y-%m-%d")
    text = _cell_str(value)
    if len(text) >= 10 and text[4] == "-" and text[7] == "-":
        return text[:10]
    return text


def _parse_amount(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return round(float(value), 2)
    text = _cell_str(value).replace(",", "")
    try:
        return round(float(text), 2)
    except ValueError:
        return None


def _row_is_header(row: tuple[Any, ...]) -> bool:
    if not row:
        return False
    return _cell_str(row[0]) == HEADER_MARKER


def _normalize_row(row: tuple[Any, ...]) -> tuple[Any, ...]:
    row_list = list(row)
    while len(row_list) < _SHEET_COLUMN_COUNT:
        row_list.append(None)
    return tuple(row_list[:_SHEET_COLUMN_COUNT])


def _parse_transaction_row(row: tuple[Any, ...]) -> Transaction | None:
    row = _normalize_row(row)
    if all(v is None or _cell_str(v) == "" for v in row):
        return None

    date_str = _excel_date(row[0])
    tx_type = _cell_str(row[1])
    if not date_str and not tx_type:
        return None
    if _cell_str(row[0]) == HEADER_MARKER:
        return None

    return Transaction(
        交易时间=date_str,
        交易类型=tx_type,
        交易对方=_cell_str(row[2]),
        收_支=_cell_str(row[4]),
        金额=_parse_amount(row[5]),
    )


def iter_sheet_rows(path: str | Path, sheet_name: str | None = None) -> Iterator[tuple[Any, ...]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet_name] if sheet_name else wb.active
        for row in ws.iter_rows(values_only=True):
            yield row
    finally:
        wb.close()


def parse_wechat_bill(path: str | Path, sheet_name: str | None = None) -> list[Transaction]:
    """解析微信支付 Excel 账单，返回交易列表。"""
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(path)

    rows = list(iter_sheet_rows(path, sheet_name))
    header_idx: int | None = None

    for i, row in enumerate(rows):
        if _row_is_header(row):
            header_idx = i
            break

    if header_idx is None:
        raise ValueError(
            f"未在文件中找到明细表头（需包含「{HEADER_MARKER}」）。"
            f"请确认是否为微信支付导出的 xlsx 账单。"
        )

    transactions: list[Transaction] = []
    for row in rows[header_idx + 1 :]:
        tx = _parse_transaction_row(row)
        if tx is not None:
            transactions.append(tx)

    return transactions


def main() -> None:
    parser = argparse.ArgumentParser(description="解析微信支付 Excel 账单")
    parser.add_argument(
        "xlsx_path",
        nargs="?",
        default=r"d:\Downloads\微信支付账单流水文件(20260821-20260921)_20260921185618.xlsx",
        help="账单 xlsx 路径",
    )
    parser.add_argument(
        "-o",
        "--output",
        help="将 JSON 结果写入该文件（UTF-8）",
    )
    parser.add_argument(
        "--limit",
        type=int,
        default=0,
        help="仅打印前 N 笔（0 表示全部）",
    )
    args = parser.parse_args()

    transactions = parse_wechat_bill(args.xlsx_path)
    records = [t.to_dict() for t in transactions]

    if args.output:
        out = Path(args.output)
        out.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"已写入 {out}，共 {len(transactions)} 笔")
        return

    print(f"共 {len(transactions)} 笔\n")
    show = transactions if args.limit <= 0 else transactions[: args.limit]
    for i, tx in enumerate(show, 1):
        d = tx.to_dict()
        print(f"--- #{i} ---")
        for k, v in d.items():
            print(f"  {k}: {v}")


if __name__ == "__main__":
    main()
